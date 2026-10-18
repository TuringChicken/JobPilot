import re
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

SUPPORTED_SUFFIXES = {".pdf", ".docx"}
MIN_TEXT_CHARS = 50


class UnsupportedResumeType(ValueError):
    pass


class ResumeExtractionError(ValueError):
    pass


def extract_resume_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise UnsupportedResumeType(f"Unsupported file type '{suffix}'. Upload a PDF or DOCX.")

    # Check magic bytes so a renamed file can't slip through on its extension alone.
    if suffix == ".pdf":
        if not content.startswith(b"%PDF"):
            raise ResumeExtractionError("File is not a valid PDF.")
        text = _extract_pdf(content)
    else:
        if not content.startswith(b"PK\x03\x04"):
            raise ResumeExtractionError("File is not a valid DOCX.")
        text = _extract_docx(content)

    text = _normalize(text)
    if len(text) < MIN_TEXT_CHARS:
        raise ResumeExtractionError(
            "Could not extract enough text. Scanned or image-only resumes are not supported yet."
        )
    return text


def _extract_pdf(content: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted:
            raise ResumeExtractionError("Password-protected PDFs are not supported.")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except PdfReadError as exc:
        raise ResumeExtractionError("Could not read PDF.") from exc


def _extract_docx(content: bytes) -> str:
    try:
        doc = Document(BytesIO(content))
    except Exception as exc:  # python-docx raises several unrelated types for corrupt files
        raise ResumeExtractionError("Could not read DOCX.") from exc

    parts = [p.text for p in doc.paragraphs]
    # Skills are often laid out in tables, which doc.paragraphs does not include.
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _normalize(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()
