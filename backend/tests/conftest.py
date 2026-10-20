from collections.abc import Iterator
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.api.deps import get_job_parser, get_match_evaluator, get_resume_parser
from app.config import Settings
from app.db import Base
from app.main import create_app
from app.schemas.jobs import (
    DimensionAssessment,
    Evidence,
    JobPostingData,
    MatchEvaluation,
    SkillGap,
    SkillMatch,
)
from app.schemas.profile import CandidateProfileData, Skill

API_KEY = "test-key"
SAMPLE_RESUME = (
    "Software developer with three years of Python experience. "
    "Built REST APIs with FastAPI and deployed services on Azure. "
    "Bachelor of Computer Science, University of Auckland."
)


class FakeResumeParser:
    def __init__(self) -> None:
        self.calls: list[str] = []

    def parse(self, resume_text: str) -> CandidateProfileData:
        self.calls.append(resume_text)
        return CandidateProfileData(
            summary="Python developer",
            work_rights=None,
            education=[],
            skills=[Skill(name="Python", category="language", proficiency=None)],
            certifications=[],
            languages=[],
            awards=[],
            work_experience=[],
            projects=[],
        )


SAMPLE_JOB_TEXT = (
    "Junior AI Engineer at Example Ltd, Auckland. Hybrid, full-time. "
    "Salary $75,000 - $85,000 per year. "
    "You will build LLM-powered features and REST APIs in Python. "
    "Required: Python, FastAPI, Azure. Nice to have: Docker. "
    "Visa sponsorship available for the right candidate."
)


def make_job_data(**overrides) -> JobPostingData:
    fields = dict(
        is_job_posting=True,
        title="Junior AI Engineer",
        company="Example Ltd",
        location="Auckland",
        work_arrangement="hybrid",
        employment_type="full_time",
        experience_level="junior",
        min_years_experience=None,
        salary_min=75000,
        salary_max=85000,
        salary_currency="NZD",
        salary_period="year",
        required_skills=["Python", "FastAPI", "Azure"],
        preferred_skills=["Docker"],
        responsibilities=["Build LLM-powered features"],
        qualifications=[],
        work_rights=None,
    )
    return JobPostingData(**(fields | overrides))


def make_evaluation(**overrides) -> MatchEvaluation:
    def dim(score: int) -> DimensionAssessment:
        return DimensionAssessment(
            score=score,
            rationale="Fake rationale",
            evidence=[
                Evidence(source="cv", quote="Python developer"),
                Evidence(source="jd", quote="REST APIs in Python"),
            ],
        )

    fields = dict(
        skills=dim(80),
        experience=dim(60),
        responsibilities=dim(70),
        matched_skills=[
            SkillMatch(skill="Python", evidence=[Evidence(source="cv", quote="Python")])
        ],
        missing_skills=[SkillGap(skill="Azure", importance="required", note=None)],
        condition_checks=[],
        uncertainties=["Work rights not stated"],
    )
    return MatchEvaluation(**(fields | overrides))


class FakeJobParser:
    def __init__(self) -> None:
        self.calls: list[str] = []
        self.result = make_job_data()

    def parse(self, job_text: str) -> JobPostingData:
        self.calls.append(job_text)
        return self.result


class FakeMatchEvaluator:
    def __init__(self) -> None:
        self.calls: list[dict] = []
        self.result = make_evaluation()

    def evaluate(self, profile, preferences, job, job_text, language="en") -> MatchEvaluation:
        self.calls.append(
            {
                "profile": profile,
                "preferences": preferences,
                "job": job,
                "job_text": job_text,
                "language": language,
            }
        )
        return self.result


@pytest.fixture(autouse=True)
def _isolate_llm_env(monkeypatch):
    """Keep credentials from the developer's shell out of tests, so nothing hits a real API."""
    for name in Settings.model_fields:
        monkeypatch.delenv(name.upper(), raising=False)


@pytest.fixture
def settings(tmp_path) -> Settings:
    return Settings(
        _env_file=None,
        api_key=API_KEY,
        # No LLM credentials => no real LLM client; tests inject a fake parser.
        database_url=f"sqlite:///{tmp_path / 'test.db'}",
        max_resume_bytes=200_000,
    )


@pytest.fixture
def fake_parser() -> FakeResumeParser:
    return FakeResumeParser()


@pytest.fixture
def fake_job_parser() -> FakeJobParser:
    return FakeJobParser()


@pytest.fixture
def fake_evaluator() -> FakeMatchEvaluator:
    return FakeMatchEvaluator()


@pytest.fixture
def client(settings, fake_parser, fake_job_parser, fake_evaluator) -> Iterator[TestClient]:
    app = create_app(settings)
    Base.metadata.create_all(app.state.engine)
    app.dependency_overrides[get_resume_parser] = lambda: fake_parser
    app.dependency_overrides[get_job_parser] = lambda: fake_job_parser
    app.dependency_overrides[get_match_evaluator] = lambda: fake_evaluator
    with TestClient(app, headers={"X-API-Key": API_KEY}) as c:
        yield c


def make_docx(text: str = SAMPLE_RESUME) -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "Docker, PostgreSQL"
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_pdf(text: str = SAMPLE_RESUME) -> bytes:
    """Build a minimal single-page PDF with real text, with a correct xref table."""
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 10 Tf 40 750 Td ({escaped}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref_pos = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref_pos,
    )
    return bytes(out)
